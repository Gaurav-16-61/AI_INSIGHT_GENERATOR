"""Universal file router: tabular / text-documents / images / anything else."""
import base64, hashlib, html, io, json, os, re, zipfile
from collections import Counter
from datetime import datetime
from html.parser import HTMLParser

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

import insight_engine as ie

TABULAR = {"csv", "tsv", "xlsx", "xls", "json"}
IMAGE = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "tiff", "tif"}
MEDIA = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp"}
STOP = set("""the a an and or but if of to in on at by for with from as is are was were be been being it its this that these
those i you he she we they them his her our your their not no so do does did have has had will would can could should may might
than then there here what which who whom when where why how into over under about after before also just more most other some
such only own same too very am up out all any both each few one two per via etc""".split())
POS = set("good great excellent amazing love loved best happy wonderful fantastic nice awesome perfect positive easy fast helpful recommend enjoy enjoyed success improved better".split())
NEG = set("bad terrible awful worst hate hated poor slow broken bug error fail failed failure problem issue issues difficult negative disappointing disappointed waste useless crash crashed complaint refund delay delayed".split())
MAGIC = {b"%PDF": "PDF", b"PK\x03\x04": "ZIP-based archive", b"\x89PNG": "PNG image", b"\xff\xd8\xff": "JPEG image",
         b"GIF8": "GIF image", b"MZ": "Windows executable", b"\x7fELF": "Linux executable", b"SQLite format 3": "SQLite database",
         b"ID3": "MP3 audio", b"\x1f\x8b": "GZIP archive", b"Rar!": "RAR archive", b"7z\xbc\xaf": "7-Zip archive"}


def _b64(fig):
    return ie._fig_to_b64(fig)


def _bar(labels, values, title, color="#3b6fd4", horizontal=True):
    fig, ax = plt.subplots(figsize=(5, 3))
    if horizontal:
        ax.barh(labels[::-1], values[::-1], color=color)
    else:
        ax.bar(labels, values, color=color)
        plt.xticks(rotation=45, ha="right")
    ax.set_title(title)
    return _b64(fig)


# ---------------- tabular ----------------
def load_tabular(data, ext):
    bio = io.BytesIO(data)
    if ext == "csv":
        df = pd.read_csv(bio, sep=None, engine="python")
    elif ext == "tsv":
        df = pd.read_csv(bio, sep="\t")
    elif ext in ("xlsx", "xls"):
        df = pd.read_excel(bio)
    else:
        j = json.loads(data.decode("utf-8", "ignore"))
        if isinstance(j, dict):
            j = next((v for v in j.values() if isinstance(v, list) and v and isinstance(v[0], dict)), [j])
        df = pd.json_normalize(j)
    if df.shape[1] < 2 and ext in ("csv", "tsv") or df.empty:
        raise ValueError("not a real table")
    return ie.coerce_dtypes(df)


def analyze_tabular(data, ext):
    df = load_tabular(data, ext)
    p = ie.profile(df)
    ins = ie.generate_insights(df, p)
    charts = ie.make_charts(df, p)
    for c in p["categorical"]:  # free-text columns get text analysis too
        s = df[c].dropna().astype(str)
        if len(s) > 5 and s.str.len().mean() > 40:
            t = analyze_text("\n".join(s.head(5000)), name=c)
            ins.append({"severity": "info", "category": "Text column", "text": f"'{c}' contains free text. " + t["insights"][0]["text"] if t["insights"] else ""})
            charts.update({f"[{c}] {k}": v for k, v in t["charts"].items() if k.startswith("Top words")})
            break
    kpis = [("Rows", f"{p['rows']:,}"), ("Columns", p["columns"]), ("Duplicates", p["duplicates"]), ("Insights", len(ins))]
    return {"kind": "table", "kind_label": "Structured table", "kpis": kpis, "insights": ins, "charts": charts,
            "columns": p["columns_info"], "preview_html": df.head(10).to_html(index=False, border=0),
            "facts": {"rows": p["rows"], "columns": p["columns_info"][:40]}}


# ---------------- text ----------------
class _Strip(HTMLParser):
    def __init__(self):
        super().__init__(); self.out = []; self.skip = 0
    def handle_starttag(self, t, a):
        if t in ("script", "style"): self.skip += 1
    def handle_endtag(self, t):
        if t in ("script", "style"): self.skip = max(0, self.skip - 1)
    def handle_data(self, d):
        if not self.skip: self.out.append(d)


def extract_text(data, ext):
    """Returns (text, meta) or (None, {}) when the file is not readable text."""
    if ext == "pdf":
        from pypdf import PdfReader
        rd = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in rd.pages), {"Pages": len(rd.pages)}
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs] + [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
        return "\n".join(parts), {"Tables": len(d.tables)}
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame), {"Slides": len(prs.slides)}
    if b"\x00" in data[:4096]:
        return None, {}
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    if ext in ("html", "htm", "xml"):
        s = _Strip(); s.feed(text); text = " ".join(s.out)
    return text, {}


def analyze_text(text, meta=None, name=""):
    text = text[:2_000_000]
    lw = [w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'’-]*", text)]
    nw = len(lw)
    sents = [s for s in re.split(r"(?<=[.!?])\s+|\n+", text) if s.strip()]
    lens = [len(s.split()) for s in sents]
    content = [w for w in lw if w not in STOP and len(w) > 2]
    top = Counter(content).most_common(10)
    bi = [(" ".join(k), v) for k, v in Counter(zip(content, content[1:])).most_common(8) if v > 1]
    syl = sum(max(1, len(re.findall(r"[aeiouy]+", w))) for w in lw)
    avg_len = nw / max(1, len(sents))
    ease = 206.835 - 1.015 * avg_len - 84.6 * (syl / max(1, nw))
    pos, neg = sum(w in POS for w in lw), sum(w in NEG for w in lw)
    ins = []
    add = lambda s, c, t: ins.append({"severity": s, "category": c, "text": t})
    if top:
        add("info", "Themes", "Most frequent topics: " + ", ".join(w for w, _ in top[:6]) + ".")
    if bi:
        add("info", "Phrases", "Recurring phrases: " + ", ".join(f"'{k}' ({v}x)" for k, v in bi[:4]) + ".")
    if nw < 50:
        add("warning", "Length", f"Only {nw} words - too little text for reliable statistics.")
    if nw > 30:
        band = "very easy" if ease > 80 else "fairly easy" if ease > 60 else "moderate" if ease > 40 else "difficult" if ease > 20 else "very difficult"
        add("info", "Readability", f"Reading ease is {ease:.0f}/100 ({band}); average sentence is {avg_len:.1f} words.")
    if avg_len > 25 and nw > 100:
        add("warning", "Readability", "Sentences are very long on average - consider shortening them.")
    if pos + neg >= 3:
        tone = "positive" if pos > neg * 1.5 else "negative" if neg > pos * 1.5 else "mixed/neutral"
        add("info", "Sentiment", f"Overall tone looks {tone} ({pos} positive vs {neg} negative cue words).")
    if nw > 100 and len(set(lw)) / nw < 0.25:
        add("warning", "Repetition", f"Low vocabulary variety ({len(set(lw))/nw:.0%} unique words) - text is highly repetitive.")
    lines = [l.strip() for l in text.splitlines() if len(l.strip()) > 15]
    dl = Counter(lines).most_common(1)
    if dl and dl[0][1] >= 3:
        add("warning", "Repetition", f"One line repeats {dl[0][1]} times: '{dl[0][0][:80]}'.")
    levels = Counter(re.findall(r"\b(CRITICAL|FATAL|ERROR|WARN(?:ING)?|INFO|DEBUG)\b", text))
    if sum(levels.values()) >= 5:
        e = levels["ERROR"] + levels["CRITICAL"] + levels["FATAL"]
        add("critical" if e else "info", "Log analysis", "Looks like a log file: " + ", ".join(f"{k}={v}" for k, v in levels.most_common()) + ".")
    emails, urls = len(set(re.findall(r"[\w.+-]+@[\w-]+\.[\w.]+", text))), len(set(re.findall(r"https?://\S+", text)))
    if emails:
        add("warning", "Privacy", f"Contains {emails} distinct email address(es) - possible personal data.")
    if urls:
        add("info", "Links", f"Contains {urls} distinct URL(s).")
    charts = {}
    if top:
        charts["Top words"] = _bar([w for w, _ in top], [v for _, v in top], "Top words")
    if bi:
        charts["Top phrases"] = _bar([k for k, _ in bi], [v for _, v in bi], "Top phrases", "#3ba07a")
    if len(lens) > 5:
        fig, ax = plt.subplots(figsize=(5, 3)); ax.hist([min(l, 60) for l in lens], bins=20, color="#e8a13a"); ax.set_title("Sentence length (words)")
        charts["Sentence lengths"] = _b64(fig)
    kpis = [("Words", f"{nw:,}"), ("Sentences", f"{len(sents):,}"), ("Unique words", f"{len(set(lw)):,}"), ("Reading ease", f"{ease:.0f}" if nw > 30 else "-")]
    kpis += list((meta or {}).items())
    preview = "<pre style='white-space:pre-wrap;margin:0'>" + html.escape(text[:1500]) + ("…" if len(text) > 1500 else "") + "</pre>"
    return {"kind": "text", "kind_label": "Text document", "kpis": kpis, "insights": ins, "charts": charts, "columns": None,
            "preview_html": preview, "facts": {"words": nw, "sentences": len(sents), "top_words": top, "sentiment": [pos, neg]},
            "excerpt": text[:8000]}


# ---------------- images ----------------
def analyze_image(data, ext):
    from PIL import Image
    im = Image.open(io.BytesIO(data))
    fmt, (w, h), mode = im.format, im.size, im.mode
    frames = getattr(im, "n_frames", 1)
    small = im.convert("RGB"); small.thumbnail((256, 256))
    arr = np.asarray(small).astype(float); gray = arr.mean(axis=2)
    bright, contrast = gray.mean(), gray.std()
    gy, gx = np.gradient(gray); sharp = float(np.var(gx) + np.var(gy))
    q = small.quantize(colors=5); pal = q.getpalette()[:15]
    cols = sorted(q.getcolors(), reverse=True)
    dom = [("#%02x%02x%02x" % tuple(pal[i * 3:i * 3 + 3]), n / sum(c for c, _ in cols)) for n, i in cols]
    ins = []
    add = lambda s, c, t: ins.append({"severity": s, "category": c, "text": t})
    add("info", "Format", f"{fmt} image, {w}x{h}px ({w*h/1e6:.1f} MP), mode {mode}" + (f", {frames} frames (animated)" if frames > 1 else "") + ".")
    if min(w, h) < 300: add("warning", "Quality", "Low resolution - fine detail will be limited.")
    if bright < 60: add("warning", "Exposure", f"Image is dark (mean brightness {bright:.0f}/255).")
    elif bright > 200: add("warning", "Exposure", f"Image is very bright (mean brightness {bright:.0f}/255).")
    if contrast < 30: add("warning", "Contrast", "Low contrast - the image looks flat or washed out.")
    if sharp < 15: add("warning", "Sharpness", "Image may be blurry or very smooth (low edge detail).")
    add("info", "Colours", "Dominant colours: " + ", ".join(f"{c} ({p:.0%})" for c, p in dom[:3]) + ".")
    try:
        ex = im.getexif()
        if ex.get_ifd(0x8825):
            add("critical", "Privacy", "Contains GPS location metadata - remove it before sharing publicly.")
        if ex.get(271) or ex.get(272):
            add("info", "Camera", f"Taken with {ex.get(271, '')} {ex.get(272, '')}".strip() + ".")
    except Exception:
        pass
    charts = {}
    fig, ax = plt.subplots(figsize=(5, 1.6))
    ax.barh([0], [p for _, p in dom], color=[c for c, _ in dom], left=np.cumsum([0] + [p for _, p in dom[:-1]]))
    ax.axis("off"); ax.set_title("Colour palette"); charts["Colour palette"] = _b64(fig)
    fig, ax = plt.subplots(figsize=(5, 3))
    for i, c in enumerate("rgb"): ax.hist(arr[..., i].ravel(), bins=32, alpha=.5, color=c)
    ax.set_title("RGB histogram"); charts["RGB histogram"] = _b64(fig)
    kpis = [("Width", w), ("Height", h), ("Brightness", f"{bright:.0f}"), ("Contrast", f"{contrast:.0f}")]
    return {"kind": "image", "kind_label": "Image", "kpis": kpis, "insights": ins, "charts": charts, "columns": None,
            "preview_html": None, "facts": {"format": fmt, "size": [w, h], "brightness": bright, "contrast": contrast, "colours": dom[:3]},
            "_image": (data, MEDIA.get(ext))}


# ---------------- anything else ----------------
def analyze_other(data, ext, note=""):
    n = len(data)
    magic = next((v for k, v in MAGIC.items() if data.startswith(k)), "unknown binary")
    counts = np.bincount(np.frombuffer(data[:2_000_000], dtype=np.uint8), minlength=256)
    pr = counts[counts > 0] / counts.sum() if counts.sum() else np.array([1.0])
    ent = float(-(pr * np.log2(pr)).sum())
    ins = [{"severity": "info", "category": "File", "text": f"{n/1024:,.1f} KB, detected as {magic} (extension: .{ext or 'none'})."},
           {"severity": "warning", "category": "Limits", "text": note or "This file's content can't be read as a table, text or image, so only file-level facts are shown."}]
    if ent > 7.5: ins.append({"severity": "info", "category": "Entropy", "text": f"Very high entropy ({ent:.2f}/8) - data is compressed or encrypted."})
    charts, kpis = {}, [("Size (KB)", f"{n/1024:,.1f}"), ("Type", magic), ("Entropy", f"{ent:.2f}"), ("SHA-256", hashlib.sha256(data).hexdigest()[:10])]
    if zipfile.is_zipfile(io.BytesIO(data)):
        z = zipfile.ZipFile(io.BytesIO(data)); infos = [i for i in z.infolist() if not i.is_dir()]
        ext_c = Counter(os.path.splitext(i.filename)[1].lower() or "(none)" for i in infos).most_common(8)
        ins.append({"severity": "info", "category": "Archive", "text": f"Contains {len(infos)} files, {sum(i.file_size for i in infos)/1e6:.1f} MB uncompressed. Types: " + ", ".join(f"{k}={v}" for k, v in ext_c) + "."})
        if ext_c: charts["Files by type"] = _bar([k for k, _ in ext_c], [v for _, v in ext_c], "Files by type")
    fig, ax = plt.subplots(figsize=(5, 3)); ax.bar(range(256), counts, color="#6b7488"); ax.set_title("Byte value distribution")
    charts["Byte distribution"] = _b64(fig)
    return {"kind": "other", "kind_label": "Other file", "kpis": kpis, "insights": ins, "charts": charts, "columns": None,
            "preview_html": None, "facts": {"size": n, "type": magic, "entropy": ent}}


# ---------------- router ----------------
def analyze_file(data: bytes, filename: str) -> dict:
    if not data:
        raise ValueError("The file is empty.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    note = ""
    if ext in TABULAR:
        try:
            return analyze_tabular(data, ext)
        except Exception as e:
            note = f"Could not read as a table ({e}); analysed as text instead."
    if ext in IMAGE:
        try:
            return analyze_image(data, ext)
        except Exception as e:
            note = f"Could not decode image: {e}"
    else:
        try:
            text, meta = extract_text(data, ext)
        except Exception as e:
            text, meta, note = None, {}, f"Could not read .{ext} content: {e}"
        if text and len(re.findall(r"\w+", text)) > 0:
            res = analyze_text(text, meta)
            if note: res["insights"].insert(0, {"severity": "info", "category": "Note", "text": note})
            return res
    return analyze_other(data, ext, note)


# ---------------- AI summary ----------------
def ai_summary(res):
    if not os.getenv("ANTHROPIC_API_KEY"):
        return None
    try:
        import anthropic
        bullets = "\n".join(f"- {i['text']}" for i in res["insights"][:40])
        ask = {"table": "an executive summary, key data risks, cleaning steps, and suggested analysis/ML directions",
               "text": "a summary of the content, main themes, tone, notable concerns, and suggested next steps",
               "image": "what the image shows, its quality issues, and suggestions for improvement or use",
               "other": "what this file probably is, and what the user should do to analyse it properly"}[res["kind"]]
        prompt = (f"You are a senior data analyst. File: {res['filename']} ({res['kind_label']}).\nFacts: {json.dumps(res['facts'], default=str)[:6000]}\n"
                  f"Automated findings:\n{bullets}\n" + (f"\nContent excerpt:\n{res['excerpt']}\n" if res.get("excerpt") else "") +
                  f"\nWrite {ask}. Use only the information given; be specific, no markdown headers.")
        content = [{"type": "text", "text": prompt}]
        img = res.get("_image")
        if img and img[1] and len(img[0]) < 4_500_000:
            content.insert(0, {"type": "image", "source": {"type": "base64", "media_type": img[1], "data": base64.b64encode(img[0]).decode()}})
        msg = anthropic.Anthropic().messages.create(model=os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5-5"), max_tokens=1200,
                                                    messages=[{"role": "user", "content": content}])
        return msg.content[0].text
    except Exception as e:
        return f"(AI summary unavailable: {e})"


# ---------------- report ----------------
def build_report(res, narrative=None):
    e = html.escape
    color = {"critical": "#d9534f", "warning": "#e8a13a", "info": "#3ba0c9"}
    kp = "".join(f"<div><b>{e(str(v))}</b>{e(str(k))}</div>" for k, v in res["kpis"])
    ins = "".join(f"<li><span class='b' style='background:{color[i['severity']]}'>{i['severity']}</span><b>{e(i['category'])}:</b> {e(i['text'])}</li>" for i in res["insights"]) or "<li>No issues found.</li>"
    imgs = "".join(f"<img src='data:image/png;base64,{b}' alt='{e(t)}'>" for t, b in res["charts"].items())
    cols = ""
    if res.get("columns"):
        rows = "".join(f"<tr><td>{e(str(c['name']))}</td><td>{c['dtype']}</td><td>{c['missing_pct']}%</td><td>{c['unique']}</td><td>{c.get('mean','')}</td><td>{c.get('min','')}</td><td>{c.get('max','')}</td></tr>" for c in res["columns"])
        cols = f"<h2>Column summary</h2><table><tr><th>Column</th><th>Type</th><th>Missing</th><th>Unique</th><th>Mean</th><th>Min</th><th>Max</th></tr>{rows}</table>"
    ai = f"<h2>AI Summary</h2><div class='n'>{e(narrative)}</div>" if narrative else ""
    pv = f"<h2>Preview</h2>{res['preview_html']}" if res.get("preview_html") else ""
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>Insight Report - {e(res['filename'])}</title><style>
body{{font-family:system-ui,sans-serif;max-width:1000px;margin:2rem auto;padding:0 1rem;color:#222}}h1{{border-bottom:3px solid #3b6fd4}}
.k{{display:flex;gap:1rem;flex-wrap:wrap}}.k div{{background:#f3f6fb;padding:1rem;border-radius:8px;flex:1;min-width:110px;text-align:center}}.k b{{font-size:1.4rem;display:block}}
.b{{color:#fff;padding:2px 8px;border-radius:10px;font-size:.7rem;margin-right:8px;text-transform:uppercase}}li{{margin:.5rem 0}}
table{{border-collapse:collapse;font-size:.85rem}}td,th{{padding:5px 9px;border-bottom:1px solid #ddd;text-align:left}}img{{width:470px;max-width:100%;margin:6px}}
.n{{white-space:pre-wrap;background:#f9f9f9;padding:1rem;border-left:4px solid #3b6fd4}}</style></head><body>
<h1>AI Data Science Insight Report</h1><p>File: <b>{e(res['filename'])}</b> &middot; {e(res['kind_label'])} &middot; {datetime.now():%Y-%m-%d %H:%M}</p>
<div class="k">{kp}</div>{ai}<h2>Key insights</h2><ul>{ins}</ul>{cols}<h2>Visualizations</h2>{imgs}{pv}</body></html>"""
